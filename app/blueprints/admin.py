import cv2
import pandas as pd
import io
from datetime import datetime, timedelta
from flask import Blueprint, render_template, session, Response, request, redirect, url_for, jsonify, send_file
from app.services.embedder import extract_and_save_embeddings
from flask_login import login_required, logout_user
# 1. THÊM 2 DÒNG NÀY ĐỂ KẾT NỐI DATABASE VÀ MODEL SINH VIÊN
from app.models import db 
from app.models.student import Student
from app.models.attendance import Attendance
from app.services.face_detector import crop_and_save_face
from app.services.excel_service import get_attendance_dataframe, get_summary_dataframe
from app.services.face_recognizer import reload_embeddings_to_ram
from openpyxl.utils import get_column_letter

# Khai báo Blueprint admin_bp
admin_bp = Blueprint('admin', __name__)
latest_frame = None

@admin_bp.route('/')
@login_required
def dashboard():
    return render_template('admin/dashboard.html')

@admin_bp.route('/add')
@login_required
def add_student():
    return render_template('admin/add_student.html')

@admin_bp.route('/export')
@login_required
def export_page():
    classes_query = db.session.query(Student.class_name).distinct().all()
    class_list = [c[0] for c in classes_query if c[0]] 

    now_vn = datetime.utcnow() + timedelta(hours=7)
    today_str = now_vn.strftime('%Y-%m-%d')
    first_day_str = now_vn.replace(day=1).strftime('%Y-%m-%d')
    history = session.get('export_history', [])
    return render_template('admin/export_excel.html', classes=class_list, today=today_str, first_day=first_day_str, history=history)

def create_excel_response(df, sheet_name, filename):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]
        
        # Đã sửa lỗi: Dùng get_column_letter để tự động sinh tên cột (A, B... Z, AA, AB...)
        for idx, col in enumerate(df.columns, start=1):
            max_len = max(df[col].astype(str).map(len).max(), len(col)) + 3
            col_letter = get_column_letter(idx) 
            worksheet.column_dimensions[col_letter].width = max_len
            
    output.seek(0)
    return send_file(
        output, 
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 
        as_attachment=True, 
        download_name=filename
    )

# HÀM MỚI: LƯU LỊCH SỬ XUẤT EXCEL VÀO SESSION
def save_export_history(filename, export_type_name, url):
    # Nếu chưa có lịch sử thì tạo danh sách trống
    if 'export_history' not in session:
        session['export_history'] = []
    
    history = session['export_history']
    now = datetime.utcnow() + timedelta(hours=7)
    
    # Thêm bản ghi mới vào ĐẦU danh sách
    history.insert(0, {
        'filename': filename,
        'type': export_type_name,
        'time': now.strftime('%H:%M - %d/%m/%Y'),
        'url': url
    })
    
    # Chỉ giữ lại 5 lần xuất gần nhất cho gọn giao diện
    session['export_history'] = history[:5] 
    session.modified = True

@admin_bp.route('/download_today_excel', methods=['GET'])
@login_required
def download_today_excel():
    try:
        now_utc = datetime.utcnow()
        now_vn = now_utc + timedelta(hours=7)
        start_of_day_vn = now_vn.replace(hour=0, minute=0, second=0, microsecond=0)
        end_of_day_vn = now_vn.replace(hour=23, minute=59, second=59, microsecond=999999)

        df = get_attendance_dataframe(start_of_day_vn, end_of_day_vn)    

        date_str = start_of_day_vn.strftime('%d-%m-%Y')
        filename = f"DanhSachDiemDanh_{date_str}.xlsx"

        save_export_history(filename, "Theo ngày", url_for('admin.download_today_excel'))

        return create_excel_response(df, 'DiemDanh_HomNay', filename)
    
    except Exception as e:
        print(f"Lỗi tạo Excel hôm nay: {e}")
        return jsonify({"status": "error", "message": "Không thể xuất file Excel."}), 500

@admin_bp.route('/process_export', methods=['GET'])
@login_required
def process_export():
    try:
        export_type = request.args.get('export_type')
        
        if export_type == 'daily':
            date_str = request.args.get('target_date')
            class_name = request.args.get('class_name')
            target_date = datetime.strptime(date_str, '%Y-%m-%d')
            start = target_date.replace(hour=0, minute=0, second=0)
            end = target_date.replace(hour=23, minute=59, second=59)
            
            df = get_attendance_dataframe(start, end, class_name)
            if class_name and class_name != 'all':
                filename = f"DiemDanh_{class_name}_{target_date.strftime('%d-%m-%Y')}.xlsx"
            else:
                filename = f"DiemDanh_TatCa_{target_date.strftime('%d-%m-%Y')}.xlsx"

            save_export_history(filename, "Theo ngày", request.url)
            return create_excel_response(df, "BaoCao_Ngay", filename)

        elif export_type == 'summary':
            start_str = request.args.get('start_date')
            end_str = request.args.get('end_date')
            
            start = datetime.strptime(start_str, '%Y-%m-%d').replace(hour=0, minute=0, second=0)
            end = datetime.strptime(end_str, '%Y-%m-%d').replace(hour=23, minute=59, second=59)
            
            df = get_summary_dataframe(start, end)
            filename = f"TongHop_{start.strftime('%d%m')}_den_{end.strftime('%d%m%Y')}.xlsx"
            save_export_history(filename, "Tổng hợp", request.url)
            return create_excel_response(df, "BaoCao_TongHop", filename)
            
        else:
            return jsonify({"status": "error", "message": "Loại xuất file không hợp lệ"}), 400

    except Exception as e:
        print(f"Lỗi hệ thống Export Excel: {e}")
        return jsonify({"status": "error", "message": "Đã xảy ra lỗi khi tạo file Excel."}), 500

@admin_bp.route('/delete_attendance/<int:log_id>', methods=['DELETE'])
def detele_attendance(log_id):
    try:
        log = Attendance.query.get(log_id)
        if not log:
            return jsonify({"status": "error", "message": "Không tìm thấy bản ghi"}), 404
        
        db.session.delete(log)
        db.session.commit()

        return jsonify({"status": "success", "message": "Đã xóa thành công"})
    except Exception as e:
        db.session.rollback()
        print(f"Lỗi xóa log điểm danh: {e}")
        return jsonify({"status": "error", "message": "Lỗi server"}), 500
    
@admin_bp.route('/logout')
@login_required
def logout():
    logout_user()  
    return redirect(url_for('auth.login'))

def generate_frames():
    global latest_frame
    camera = cv2.VideoCapture(0)

    while True:
        success, frame =  camera.read()
        if not success:
            break
        else:
            frame = cv2.flip(frame, 1)
            latest_frame = frame.copy()
            ret, buffer = cv2.imencode('.jpg', frame)
            frame = buffer.tobytes()

            yield (b'--frame\r\n'
                   b'Content-Type: image/jpeg\r\n\r\n' +frame+ b'\r\n')

@admin_bp.route('/video_feed')
def video_feed():
    return Response(generate_frames(), mimetype='multipart/x-mixed-replace; boundary=frame')

@admin_bp.route('/capture',methods=['POST'])
def capture_face():
    global latest_frame
    if latest_frame is None:
        return jsonify({"status": "error", "message": "Camera chưa bật hoặc chưa sẵn sàng!"})
    data = request.json
    student_id = data.get('student_id')
    count = data.get('count')

    # Gọi AI cắt ảnh
    success, msg = crop_and_save_face(latest_frame, student_id, count)

    if success:
        return jsonify({"status": "success", "message": msg})
    else:
        return jsonify({"status": "error", "message": msg})
    
@admin_bp.route('/add_student_db', methods=['POST'])
def add_student_db():
    data = request.json
    s_id = data.get('student_id')
    name = data.get('full_name')
    c_name = data.get('class_name')
    has_face = data.get('face_saved', False)

    if not s_id or not name or not c_name:
        return jsonify({"status": "error", "message": "Vui lòng nhập đầy đủ thông tin!"})
    

    exist_student = Student.query.filter_by(student_id=s_id).first()
    if exist_student:
        return jsonify({"status": "error", "message": "Mã số sinh viên này đã tồn tại!"})

    try:
        new_student = Student(
            student_id=s_id, 
            full_name=name, 
            class_name=c_name, 
            face_saved=has_face
        )
        db.session.add(new_student)
        db.session.commit()
        
        success, msg = extract_and_save_embeddings()
        if not success:
            return jsonify({"status": "error", "message": f"Lưu DB thành công nhưng lỗi trích xuất đặc trưng FaceNet: {msg}"})
        
        reload_embeddings_to_ram()

        return jsonify({"status": "success", "message": "Đã lưu sinh viên vào Database thành công!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"status": "error", "message": f"Lỗi DB: {str(e)}"})